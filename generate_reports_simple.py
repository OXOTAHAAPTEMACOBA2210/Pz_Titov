import os
import docx
from docx.shared import Pt

# Очищенный скрипт для создания отчетов
def create_report(filename, title, code_file, code_text):
    doc = docx.Document()
    doc.add_heading(title, 0)
    doc.add_paragraph("Студент: Кузьменко В.А., группа ДЭ 21-24")
    doc.add_heading("Код:", 1)
    doc.add_paragraph(code_text)
    doc.save(filename)

BASE_DIR = "."
create_report("Отчет_ПЗ2_Кузьменко.docx", "Отчет по ПЗ 2", "pz2.sql", open("pz2.sql", "r", encoding="utf-8").read())
create_report("Отчет_ПЗ3_Кузьменко.docx", "Отчет по ПЗ 3", "pz3.sql", open("pz3.sql", "r", encoding="utf-8").read())
create_report("Отчет_ПЗ4_Кузьменко.docx", "Отчет по ПЗ 4", "pz4.sql", open("pz4.sql", "r", encoding="utf-8").read())
create_report("Отчет_ПЗ5_Кузьменко.docx", "Отчет по ПЗ 5", "pz5_crud.py", open("pz5_crud.py", "r", encoding="utf-8").read())
create_report("Отчет_ПЗ6_Кузьменко.docx", "Отчет по ПЗ 6", "pz6_stats.py", open("pz6_stats.py", "r", encoding="utf-8").read())
create_report("Отчет_ПЗ7_Кузьменко.docx", "Отчет по ПЗ 7", "pz7_parking.sql", open("pz7_parking.sql", "r", encoding="utf-8").read())
print("Отчеты созданы!")
