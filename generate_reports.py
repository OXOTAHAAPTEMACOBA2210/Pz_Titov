import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = os.path.abspath(".")
SHOTS_DIR = os.path.join(BASE_DIR, "screenshots")

def set_style(run, bold=False, size=11):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0) # ЧЕРНЫЙ

def create_report(filename, title, code_file, code_text, screenshots=None):
    doc = docx.Document()
    
    # Заголовок
    p = doc.add_paragraph()
    run = p.add_run(title)
    set_style(run, bold=True, size=16)
    
    doc.add_paragraph("Студент: Кузьменко В.А., группа ДЭ 21-24")
    
    # Код
    doc.add_heading("Код:", 2)
    doc.add_paragraph(code_text)
    
    # Скриншоты
    if screenshots:
        doc.add_heading("Скриншоты:", 2)
        for shot in screenshots:
            path = os.path.join(SHOTS_DIR, shot)
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(6))
                
    doc.save(filename)

# Генерация
create_report("Отчет_ПЗ2_Кузьменко.docx", "Отчет по ПЗ 2", "pz2.sql", open("pz2.sql", "r", encoding="utf-8").read(), ["pz2_structure.png", "pz2_test_negative.png"])
create_report("Отчет_ПЗ3_Кузьменко.docx", "Отчет по ПЗ 3", "pz3.sql", open("pz3.sql", "r", encoding="utf-8").read(), ["pz3_schema_diagram.png", "pz3_test_fk.png"])
create_report("Отчет_ПЗ4_Кузьменко.docx", "Отчет по ПЗ 4", "pz4.sql", open("pz4.sql", "r", encoding="utf-8").read(), ["pz4_comp_data.png", "pz4_study_query.png"])
create_report("Отчет_ПЗ5_Кузьменко.docx", "Отчет по ПЗ 5", "pz5_crud.py", open("pz5_crud.py", "r", encoding="utf-8").read(), ["pz5_console_crud.png"])
create_report("Отчет_ПЗ6_Кузьменко.docx", "Отчет по ПЗ 6", "pz6_stats.py", open("pz6_stats.py", "r", encoding="utf-8").read(), ["pz6_console_stats.png"])
create_report("Отчет_ПЗ7_Кузьменко.docx", "Отчет по ПЗ 7", "pz7_parking.sql", open("pz7_parking.sql", "r", encoding="utf-8").read())
print("Отчеты созданы!")
